"""공개 가상 데이터로 로그 판정과 리포트 경로를 검증한다.

저장소 루트에서 python tools/selftest.py 로 실행한다.
"""
from __future__ import annotations

import json
import sys
import tempfile
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

BASE = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(BASE))

from core.profile import Profile
from core.runner import Runner

EXPECTED = Counter({
    ("DEMO001", "PASS"): 1,
    ("DEMO001", "FAIL"): 1,
    ("DEMO001", "TIMEOUT"): 1,
    ("DEMO002", "PASS"): 1,
    ("DEMO002", "FAIL"): 1,
    ("DEMO004", "INCOMPLETE"): 1,
    ("DEMO005", "EVENT"): 1,
})


def check(condition, message):
    if not condition:
        raise AssertionError(message)


def run() -> int:
    # 실행마다 독립된 폴더를 사용해 이전 결과와 섞이지 않게 한다.
    with tempfile.TemporaryDirectory(prefix="debug-selftest-") as tmp:
        profile = Profile(name="DEMO", rules=[str(BASE / "examples/rules.csv")],
                          all_tc=True, encoding="utf-8")
        runner = Runner(profile, out_dir=tmp, quiet=True,
                        replay=str(BASE / "examples/sample.txt"))
        check([d.tc_id for d in runner.todo] == ["DEMO006"],
              "TODO 규칙이 실행 대상에서 제외되지 않았습니다.")
        check(not runner.warnings(), "가상 규칙 사이에 충돌이 있습니다.")
        runner.run()
        check(not runner.error, f"리플레이 오류: {runner.error}")
        rows = [json.loads(line) for line in
                runner.report.jsonl_path.read_text(encoding="utf-8").splitlines()]
        actual = Counter((r["tc_id"], r["status"]) for r in rows)
        check(actual == EXPECTED,
              f"판정 불일치\n기대: {EXPECTED}\n실제: {actual}")
        event = next(r for r in rows if r["status"] == "EVENT")
        check(event["reason"] == "2회 검출", "이벤트 cooldown 집계 불일치")
        passed = next(r for r in rows if r["tc_id"] == "DEMO001"
                      and r["status"] == "PASS")
        check("id=42" in passed["evidence"] and "\x1b" not in passed["evidence"],
              "정규식 또는 ANSI 정제 결과가 리포트에 반영되지 않았습니다.")
        absent_fail = next(r for r in rows if r["tc_id"] == "DEMO002"
                           and r["status"] == "FAIL")
        check("금지 문구" in absent_fail["reason"], "ABSENT 위반 사유 누락")
        check(bool(absent_fail["evidence"]), "실패 근거 누락")

        workbook = load_workbook(runner.report.xlsx_path, read_only=True)
        try:
            check(workbook.sheetnames == ["Summary", "TC_Results"],
                  "리포트 시트 구성 불일치")
            report_rows = list(workbook["TC_Results"].iter_rows(min_row=2, values_only=True))
            expected_report = EXPECTED + Counter({("DEMO003", "NOT_EXECUTED"): 1})
            check(Counter((r[1], r[3]) for r in report_rows) == expected_report,
                  "Excel 판정 또는 NOT_EXECUTED 집계 불일치")
        finally:
            workbook.close()

        # 부팅 완료 로그가 제한 시간 이후까지 오지 않는 별도 음성 사례.
        missing = Path(tmp) / "missing.txt"
        missing.write_text("[2026-01-01 09:00:00.000] DEMO_BOOT_BEGIN\n"
                           "[2026-01-01 09:01:01.000] idle\n", encoding="utf-8")
        timeout_profile = Profile(name="ABSENT_TIMEOUT", rules=profile.rules,
                                  tc="DEMO002", encoding="utf-8")
        timeout_runner = Runner(timeout_profile, out_dir=tmp, quiet=True,
                                replay=str(missing))
        timeout_runner.run()
        check(not timeout_runner.error, f"TIMEOUT 사례 실행 오류: {timeout_runner.error}")
        check([r["status"] for r in timeout_runner.report.results] == ["TIMEOUT"],
              "ABSENT 뒤 CHECK의 TIMEOUT 판정 불일치")

    print("SELFTEST 통과: 가상 판정 7건, 이벤트 2회, TODO 제외, "
          "NOT_EXECUTED 및 JSONL/Excel 검증, 부팅 TIMEOUT")
    return 0


if __name__ == "__main__":
    try:
        sys.exit(run())
    except AssertionError as exc:
        print(f"SELFTEST 실패: {exc}", file=sys.stderr)
        sys.exit(1)
