"""
리포트 생성기.

설계 변경점
-----------
1. 판정이 나올 때마다 JSONL 한 줄을 append 한다 (빠르고 크래시에 안전).
   프로그램이 강제 종료되어도 결과는 남는다.
2. 엑셀은 주기적으로 / 종료 시에 JSONL 로부터 생성한다.
   (매 건마다 워크북 전체를 다시 만들면 결과가 쌓일수록 O(N^2) 로 느려진다)
3. 임시 파일에 쓴 뒤 os.replace 로 교체 -> 저장 중 종료돼도 파일이 깨지지 않는다.
4. 한 번도 실행되지 않은 TC 를 NOT_EXECUTED 로 반드시 집계한다.
   이게 없으면 "TC 가 아예 안 돌았는데 PASS율 100%" 라는 착시가 생긴다.
"""

from __future__ import annotations

import json
import os
from collections import Counter
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from core.verdict import EVENT, FAIL, INCOMPLETE, NOT_EXECUTED, PASS, TIMEOUT

HEADER_FILL = PatternFill("solid", start_color="305496")
HEADER_FONT = Font(name="Arial", bold=True, color="FFFFFF")
FILLS = {
    PASS: PatternFill("solid", start_color="C6EFCE"),
    FAIL: PatternFill("solid", start_color="FFC7CE"),
    TIMEOUT: PatternFill("solid", start_color="FFF2CC"),
    NOT_EXECUTED: PatternFill("solid", start_color="D9D9D9"),
    INCOMPLETE: PatternFill("solid", start_color="FCE4D6"),
    EVENT: PatternFill("solid", start_color="DDEBF7"),
}
FONTS = {
    PASS: Font(name="Arial", color="006100", bold=True),
    FAIL: Font(name="Arial", color="9C0006", bold=True),
    TIMEOUT: Font(name="Arial", color="9C6500", bold=True),
    NOT_EXECUTED: Font(name="Arial", color="595959", bold=True),
    INCOMPLETE: Font(name="Arial", color="833C00", bold=True),
    EVENT: Font(name="Arial", color="1F4E78", bold=True),
}
BASE_FONT = Font(name="Arial")
BOLD = Font(name="Arial", bold=True)
THIN = Border(*(Side(style="thin", color="BFBFBF"),) * 4)
CENTER = Alignment(horizontal="center", vertical="center")
WRAP = Alignment(wrap_text=True, vertical="top")

AUTOSAVE_EVERY = 20   # 엑셀 재생성 주기 (JSONL 은 매번 저장됨)


class ReportGenerator:
    def __init__(self, out_dir="output", product="", log_source="", rule_file="",
                 tester="", fw_version="", run_tag=""):
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        # 제품별로 하위 폴더를 나눈다 -> 리포트가 섞이지 않는다
        self.product = product
        self.out_dir = Path(out_dir) / product if product else Path(out_dir)
        self.out_dir.mkdir(parents=True, exist_ok=True)
        tagpart = f"_{product}" if product else ""
        self.xlsx_path = self.out_dir / f"{stamp}{tagpart}_QA_Report.xlsx"
        self.jsonl_path = self.out_dir / f"{stamp}{tagpart}_results.jsonl"

        self.started_at = datetime.now()
        self.log_source = str(log_source)
        self.rule_file = str(rule_file)
        self.tester = tester
        self.fw_version = fw_version
        self.run_tag = run_tag

        self.results: list[dict] = []
        self.armed: list = []          # 무장된 TCDefinition 목록
        self._unsaved = 0

    # ------------------------------------------------------------------
    def set_armed(self, definitions):
        self.armed = list(definitions)

    def add(self, verdict):
        row = {
            "tc_id": verdict.tc_id,
            "tc_name": verdict.tc_name,
            "purpose": verdict.purpose,
            "status": verdict.status,
            "start": verdict.start_time.strftime("%Y-%m-%d %H:%M:%S.%f")[:-3]
                     if verdict.start_time else "",
            "end": verdict.end_time.strftime("%Y-%m-%d %H:%M:%S.%f")[:-3]
                   if verdict.end_time else "",
            "exec_time": verdict.exec_time,
            "reason": verdict.reason,
            "failed_step": verdict.failed_step,
            "evidence": verdict.evidence_text(),
        }
        self.results.append(row)

        # 1) 즉시 JSONL append (크래시 안전)
        try:
            with open(self.jsonl_path, "a", encoding="utf-8") as fh:
                fh.write(json.dumps(row, ensure_ascii=False) + "\n")
        except OSError:
            pass

        # 2) 엑셀은 N건마다
        self._unsaved += 1
        if self._unsaved >= AUTOSAVE_EVERY:
            self.save()

    # ------------------------------------------------------------------
    def _not_executed(self) -> list[dict]:
        executed = {r["tc_id"] for r in self.results}
        rows = []
        for d in self.armed:
            if d.tc_id in executed:
                continue
            rows.append({
                "tc_id": d.tc_id, "tc_name": d.tc_name, "purpose": d.purpose,
                "status": NOT_EXECUTED, "start": "", "end": "", "exec_time": None,
                "reason": "감시 구간에서 시작 트리거가 검출되지 않음",
                "failed_step": f"STEP 1: {d.steps[0].keyword[:60]}",
                "evidence": "",
            })
        return rows

    # ------------------------------------------------------------------
    def save(self) -> Path:
        self._unsaved = 0
        rows = self.results + self._not_executed()

        wb = Workbook()
        self._sheet_summary(wb.active, rows)
        self._sheet_results(wb.create_sheet("TC_Results"), rows)

        tmp = self.xlsx_path.with_suffix(".tmp.xlsx")
        wb.save(tmp)
        os.replace(tmp, self.xlsx_path)   # 원자적 교체
        return self.xlsx_path

    # ------------------------------------------------------------------
    def _sheet_summary(self, ws, rows):
        ws.title = "Summary"
        judged = [r for r in rows if r["status"] in (PASS, FAIL, TIMEOUT, INCOMPLETE)]
        c = Counter(r["status"] for r in rows)
        n_pass, n_fail, n_to = c[PASS], c[FAIL], c[TIMEOUT]
        n_judged = len(judged)

        ws["A1"] = f"QA Automation 판정 보고서 - {self.product}" if self.product else "QA Automation 판정 보고서"
        ws["A1"].font = Font(name="Arial", bold=True, size=14)

        meta = [
            ("제품", self.product),
            ("생성 일시", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
            ("감시 시작", self.started_at.strftime("%Y-%m-%d %H:%M:%S")),
            ("로그 소스", self.log_source),
            ("규칙 파일", self.rule_file),
            ("테스터", self.tester),
            ("FW 버전", self.fw_version),
            ("실행 태그", self.run_tag),
        ]
        r = 3
        for k, v in meta:
            ws.cell(row=r, column=1, value=k).font = BOLD
            ws.cell(row=r, column=2, value=v).font = BASE_FONT
            r += 1

        r += 1
        stats = [
            ("무장 TC 수", len(self.armed)),
            ("판정 횟수", n_judged),
            ("PASS", n_pass),
            ("FAIL", n_fail),
            ("TIMEOUT", n_to),
            ("INCOMPLETE", c[INCOMPLETE]),
            ("미실행 TC", c[NOT_EXECUTED]),
            ("이벤트 감시", c[EVENT]),
            ("PASS율(%)", round(n_pass / n_judged * 100, 1) if n_judged else 0),
            ("실행률(%)",
             round((len(self.armed) - c[NOT_EXECUTED]) / len(self.armed) * 100, 1)
             if self.armed else 0),
        ]
        stat_start = r
        for k, v in stats:
            ws.cell(row=r, column=1, value=k).font = BOLD
            cell = ws.cell(row=r, column=2, value=v)
            cell.font = BASE_FONT
            if k in ("PASS", "FAIL", "TIMEOUT", "미실행 TC"):
                cell.fill = FILLS[{"PASS": PASS, "FAIL": FAIL,
                                   "TIMEOUT": TIMEOUT,
                                   "미실행 TC": NOT_EXECUTED}[k]]
            r += 1

        ws.cell(row=stat_start - 1, column=1)  # spacer

        # TC별 집계
        r += 1
        ws.cell(row=r, column=1, value="TC별 집계").font = Font(
            name="Arial", bold=True, size=12)
        r += 1
        headers = ["TC_ID", "TC_NAME", "PASS", "FAIL", "TIMEOUT", "상태", "TEST_PURPOSE"]
        for i, h in enumerate(headers, start=1):
            cell = ws.cell(row=r, column=i, value=h)
            cell.fill, cell.font, cell.border = HEADER_FILL, HEADER_FONT, THIN
        r += 1

        per = Counter((x["tc_id"], x["status"]) for x in rows)
        seen = []
        for x in rows:
            if x["tc_id"] not in seen:
                seen.append(x["tc_id"])
        info = {x["tc_id"]: x for x in rows}

        for tc_id in seen:
            p, f, t = per[(tc_id, PASS)], per[(tc_id, FAIL)], per[(tc_id, TIMEOUT)]
            if per[(tc_id, NOT_EXECUTED)]:
                state = NOT_EXECUTED
            elif per[(tc_id, EVENT)]:
                state = EVENT
            elif f or t or per[(tc_id, INCOMPLETE)]:
                state = FAIL
            else:
                state = PASS
            ws.cell(row=r, column=1, value=tc_id).font = BASE_FONT
            ws.cell(row=r, column=2, value=info[tc_id]["tc_name"]).font = BASE_FONT
            ws.cell(row=r, column=3, value=p).font = BASE_FONT
            ws.cell(row=r, column=4, value=f).font = BASE_FONT
            ws.cell(row=r, column=5, value=t).font = BASE_FONT
            sc = ws.cell(row=r, column=6, value=state)
            sc.fill, sc.font, sc.alignment = FILLS[state], FONTS[state], CENTER
            ws.cell(row=r, column=7, value=info[tc_id]["purpose"]).font = BASE_FONT
            for cidx in range(1, 8):
                ws.cell(row=r, column=cidx).border = THIN
            r += 1

        for col, w in zip("ABCDEFG", [14, 26, 8, 8, 10, 14, 52]):
            ws.column_dimensions[col].width = w

    # ------------------------------------------------------------------
    def _sheet_results(self, ws, rows):
        headers = ["No", "TC_ID", "TC_NAME", "RESULT", "START", "END",
                   "EXEC(s)", "FAILED_STEP", "REASON", "EVIDENCE"]
        ws.append(headers)
        for i in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=i)
            cell.fill, cell.font = HEADER_FILL, HEADER_FONT
            cell.alignment, cell.border = CENTER, THIN

        for n, item in enumerate(rows, start=1):
            ws.append([
                n, item["tc_id"], item["tc_name"], item["status"],
                item["start"], item["end"],
                round(item["exec_time"], 2) if item["exec_time"] is not None else "",
                item["failed_step"], item["reason"],
                (item["evidence"] or "")[:30000],
            ])
            ridx = ws.max_row
            sc = ws.cell(row=ridx, column=4)
            sc.fill = FILLS.get(item["status"], FILLS[FAIL])
            sc.font = FONTS.get(item["status"], FONTS[FAIL])
            sc.alignment = CENTER
            for cidx in range(1, len(headers) + 1):
                cell = ws.cell(row=ridx, column=cidx)
                cell.border = THIN
                if cidx == 10:
                    cell.alignment = WRAP
                elif cidx != 4:
                    cell.font = BASE_FONT

        for i, w in enumerate([5, 10, 24, 13, 21, 21, 9, 26, 40, 90], start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A2"
        if ws.max_row > 1:
            ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{ws.max_row}"
